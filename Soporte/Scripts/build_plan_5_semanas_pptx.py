# -*- coding: utf-8 -*-
"""PPT ejecutivo CVDS — mínimo texto, flujo completo, UI visual."""
from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(r"C:\pulsanet")
OUT_DIR = ROOT / "Soporte" / "Documentos"
ASSETS = OUT_DIR / "_pptx_assets_cvds_exec"
OUT_PPTX = OUT_DIR / "TACTICALPTX_CVDS_5_SEMANAS.pptx"
LOGO = ROOT / "Soporte" / "Brand" / "from-Documents" / "TacticalPtx_logo_v6_digital_cyan.png"

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF4, 0xF7, 0xF5)
PANEL2 = RGBColor(0xE8, 0xEE, 0xEB)
INK = RGBColor(0x16, 0x1A, 0x1D)
MUTED = RGBColor(0x5C, 0x66, 0x63)
GREEN_DK = RGBColor(0x00, 0x2F, 0x2A)
GOLD = RGBColor(0xA5, 0x7F, 0x2C)
LINE = RGBColor(0xC5, 0xD0, 0xCB)

META = {"apk": "1.8.46+55", "api": "1.8.21"}

WEEKS = [
    (
        "1",
        "A",
        "Análisis",
        "Se levantan los requerimientos de unidades y dependencias; se elaboran casos de uso, diccionario de datos y diagramas (flujo / entidad-relación).",
    ),
    (
        "2",
        "B",
        "Diseño",
        "Se definen las vistas a programar, consultas, reportes, pantallas de captura y los niveles de acceso del personal usuario.",
    ),
    (
        "3",
        "C",
        "Desarrollo",
        "Se programa el sistema (API, Web y Android) con base en el análisis y el diseño, generando versiones iterativas del producto.",
    ),
    (
        "4",
        "D",
        "Pruebas",
        "Se aplican pruebas por módulo y de forma global (PTT, chat, GPS, pánico) y se valida en campo antes de la entrega.",
    ),
    (
        "5",
        "E·F",
        "Impl. + Mant.",
        "Se entrega e instala en servidores, se crea el entorno de BD y se publica; el mantenimiento continúa durante la vida útil del sistema.",
    ),
]


def set_bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    sh.fill.solid()
    sh.fill.fore_color.rgb = WHITE
    sh.line.fill.background()
    spTree = slide.shapes._spTree
    sp = sh._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def txt(slide, left, top, width, height, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align
    return box


def bar(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def card(slide, left, top, width, height, fill=PANEL):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE
    sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def link(shape, target):
    try:
        shape.click_action.target_slide = target
    except Exception:
        pass


def pic(slide, path, left, top, max_w, max_h):
    path = Path(path)
    if not path.exists():
        return None
    with Image.open(path) as im:
        w, h = im.size
    wi, hi = w / 96.0, h / 96.0
    r = min(max_w / wi, max_h / hi)
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(wi * r), Inches(hi * r))


def header(slide, title=""):
    bar(slide, 0, 0, 13.333, 0.45, GREEN_DK)
    bar(slide, 0, 0.45, 13.333, 0.06, GOLD)
    if LOGO.exists():
        try:
            slide.shapes.add_picture(str(LOGO), Inches(0.25), Inches(0.05), height=Inches(0.35))
        except Exception:
            pass
    txt(slide, 1.9, 0.08, 7, 0.3, "Defensa  ·  TacticalPtx", size=12, bold=True, color=WHITE)
    if title:
        txt(slide, 8.5, 0.08, 4.5, 0.3, title, size=12, color=RGBColor(0xD0, 0xE0, 0xDC), align=PP_ALIGN.RIGHT)


def nav(slide, by, active=None):
    chips = [
        ("Inicio", "cover"),
        ("Flujo", "flujo"),
        ("S1", "s1"),
        ("S2", "s2"),
        ("S3", "s3"),
        ("S4", "s4"),
        ("S5", "s5"),
        ("UI", "ui"),
        ("ER", "er"),
    ]
    x = 0.25
    for label, key in chips:
        if key not in by:
            continue
        w = 1.15 if len(label) <= 4 else 1.35
        sh = card(slide, x, 0.65, w, 0.32, GREEN_DK if key == active else PANEL2)
        txt(slide, x, 0.68, w, 0.26, label, size=10, bold=True, color=WHITE if key == active else GREEN_DK, align=PP_ALIGN.CENTER)
        link(sh, by[key])
        x += w + 0.08


def footer(slide):
    bar(slide, 0, 7.25, 13.333, 0.25, PANEL)
    txt(slide, 0.3, 7.28, 12.7, 0.2, f"APK {META['apk']}  ·  API {META['api']}  ·  Clic en chips para navegar", size=10, color=MUTED)


def build():
    A = ASSETS
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 0 cover, 1 flujo, 2-6 weeks, 7 ui, 8 er, 9 pruebas, 10 cierre = 11 slides
    slides = [prs.slides.add_slide(blank) for _ in range(11)]
    by = {
        "cover": slides[0],
        "flujo": slides[1],
        "s1": slides[2],
        "s2": slides[3],
        "s3": slides[4],
        "s4": slides[5],
        "s5": slides[6],
        "ui": slides[7],
        "er": slides[8],
        "pruebas": slides[9],
        "cierre": slides[10],
    }

    # --- Cover ---
    s = slides[0]
    set_bg(s)
    header(s)
    txt(s, 0.6, 1.2, 12, 0.35, "CICLO DE VIDA DEL DESARROLLO DE SISTEMAS", size=13, bold=True, color=GREEN_DK)
    txt(s, 0.6, 1.7, 10, 0.8, "5 semanas", size=44, bold=True, color=INK)
    # mini flow preview
    pic(s, A / "cvds_flujo_completo.png", 0.4, 2.7, 12.5, 3.6)
    cta = card(s, 0.6, 6.5, 2.6, 0.5, GREEN_DK)
    txt(s, 0.6, 6.58, 2.6, 0.35, "Ver flujo →", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    link(cta, by["flujo"])
    cta2 = card(s, 3.5, 6.5, 2.4, 0.5, PANEL2)
    txt(s, 3.5, 6.58, 2.4, 0.35, "Semanas →", size=14, bold=True, color=GREEN_DK, align=PP_ALIGN.CENTER)
    link(cta2, by["s1"])

    # --- Flujo completo (hero) ---
    s = slides[1]
    set_bg(s)
    header(s, "Flujo completo")
    nav(s, by, "flujo")
    pic(s, A / "cvds_flujo_completo.png", 0.25, 1.15, 12.8, 5.7)
    footer(s)

    # --- Weeks (visual only) ---
    week_imgs = {
        0: ["diagrama_er.png", "diagrama_flujo_ptt.png"],
        1: ["web_login_ui.png", "web_despacho_mapa.png"],
        2: ["web_radio_chat.png", "android_radio_ptt.png"],
        3: ["pruebas_soldados_ptt.png", "android_chat.png"],
        4: ["web_despacho_mapa.png", "cvds_flujo_vertical.png"],
    }
    for i, (n, let, title, reseña) in enumerate(WEEKS):
        s = slides[2 + i]
        set_bg(s)
        header(s, f"Semana {n}")
        nav(s, by, f"s{i+1}")
        # left badge + reseña
        card(s, 0.35, 1.15, 3.55, 5.15, PANEL)
        bar(s, 0.35, 1.15, 3.55, 0.55, GREEN_DK)
        txt(s, 0.35, 1.25, 3.55, 0.4, f"S{n}  ·  Etapa {let}", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        txt(s, 0.55, 1.95, 3.15, 0.55, title, size=22, bold=True, color=INK, align=PP_ALIGN.CENTER)
        bar(s, 0.7, 2.6, 2.8, 0.05, GOLD)
        txt(s, 0.55, 2.9, 3.15, 2.8, reseña, size=13, color=MUTED, align=PP_ALIGN.LEFT)
        # images
        imgs = week_imgs[i]
        pic(s, A / imgs[0], 4.15, 1.15, 4.4, 5.15)
        pic(s, A / imgs[1], 8.75, 1.15, 4.2, 5.15)
        back = card(s, 0.35, 6.55, 1.6, 0.4, GREEN_DK)
        txt(s, 0.35, 6.6, 1.6, 0.3, "← Flujo", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        link(back, by["flujo"])
        if i < 4:
            nxt = card(s, 2.1, 6.55, 1.6, 0.4, PANEL2)
            txt(s, 2.1, 6.6, 1.6, 0.3, f"S{i+2} →", size=11, bold=True, color=GREEN_DK, align=PP_ALIGN.CENTER)
            link(nxt, slides[3 + i])
        footer(s)

    # --- UI gallery ---
    s = slides[7]
    set_bg(s)
    header(s, "UI")
    nav(s, by, "ui")
    gallery = [
        ("web_login_ui.png", 0.3, 1.15, 4.1, 2.7),
        ("web_despacho_mapa.png", 4.6, 1.15, 4.1, 2.7),
        ("web_radio_chat.png", 8.9, 1.15, 4.1, 2.7),
        ("android_radio_ptt.png", 2.2, 4.1, 4.0, 2.8),
        ("android_chat.png", 7.0, 4.1, 4.0, 2.8),
    ]
    for name, left, top, mw, mh in gallery:
        pic(s, A / name, left, top, mw, mh)
    footer(s)

    # --- ER + flows ---
    s = slides[8]
    set_bg(s)
    header(s, "Diagramas")
    nav(s, by, "er")
    pic(s, A / "diagrama_er.png", 0.25, 1.15, 8.3, 5.7)
    pic(s, A / "diagrama_flujo_ptt.png", 8.7, 1.15, 4.3, 2.7)
    pic(s, A / "diagrama_flujo_panico.png", 8.7, 4.1, 4.3, 2.7)
    footer(s)

    # --- Pruebas visual ---
    s = slides[9]
    set_bg(s)
    header(s, "Pruebas")
    nav(s, by, "pruebas")
    pic(s, A / "pruebas_soldados_ptt.png", 0.4, 1.2, 9.0, 5.6)
    for i, (label, col) in enumerate([
        ("PTT", GREEN_DK),
        ("Chat", GOLD),
        ("GPS", RGBColor(0x1E, 0x5B, 0x4F)),
        ("Pánico", RGBColor(0x9B, 0x23, 0x3C)),
    ]):
        y = 1.4 + i * 1.25
        sh = card(s, 9.7, y, 3.2, 1.0, PANEL)
        bar(s, 9.7, y, 0.18, 1.0, col)
        txt(s, 10.1, y + 0.3, 2.6, 0.4, label, size=18, bold=True, color=INK)
    footer(s)

    # --- Cierre ---
    s = slides[10]
    set_bg(s)
    header(s)
    txt(s, 0.6, 1.3, 12, 0.5, "Resumen", size=14, bold=True, color=GREEN_DK)
    txt(s, 0.6, 1.85, 12, 0.6, "Un ciclo · cinco semanas", size=32, bold=True)
    pic(s, A / "cvds_flujo_completo.png", 0.4, 2.7, 12.5, 3.5)
    home = card(s, 0.6, 6.5, 2.2, 0.45, GREEN_DK)
    txt(s, 0.6, 6.55, 2.2, 0.35, "← Flujo", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    link(home, by["flujo"])
    footer(s)

    # save
    try:
        prs.save(str(OUT_PPTX))
        print("OK", OUT_PPTX)
    except PermissionError:
        alt = OUT_DIR / "TACTICALPTX_CVDS_5_SEMANAS_v3.pptx"
        prs.save(str(alt))
        print("OK", alt)


if __name__ == "__main__":
    build()
