# -*- coding: utf-8 -*-
"""Genera FORMATO_CAMBIOS_25_09_2026.docx (Escritorio + soporte)."""
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
import shutil

SRC = r"C:\pulsanet_soporte\Documentos\FORMATO_CAMBIOS_21_09_2026.docx"
OUT = r"C:\Users\JfaRgnl_TIC\Desktop\FORMATO_CAMBIOS_25_09_2026.docx"
OUT2 = r"C:\pulsanet_soporte\Documentos\FORMATO_CAMBIOS_25_09_2026.docx"

DEV = "Sld. Inftca. Miguel Zeferino Pérez Hernández"
FECHA = "25/09/2026"

# Tipo | Descripción | Acción | Estado
CAMBIOS = [
    (
        "Avisos globales críticos.",
        "Se pueden enviar avisos a toda la fuerza (según alcance). En la app llegan a pantalla completa; Enterado detiene vibración y sonido. Se corrigió que el botón no se viera en celulares con textos largos.",
        "Módulo de avisos globales + corrección de Enterado.",
        "Aplicado",
    ),
    (
        "Renombre: pánico → Alerta.",
        "En la interfaz (web y app) el término «pánico» pasa a «Alerta», para unificar el lenguaje operativo.",
        "Cambio de nomenclatura en UI.",
        "Aplicado",
    ),
    (
        "Alcance por jerarquía (usuarios y grupos).",
        "Al alta/edición, Región → Zona → Unidad en cascada. Los admins de región/zona/unidad solo ven y asignan hacia abajo. Canales y miembros respetan el ancla geográfica.",
        "Reglas de pertenencia y alcance por rol.",
        "Aplicado",
    ),
    (
        "Perfiles: permisos por módulo y pestaña.",
        "En Administración → Perfiles se define qué módulos, pestañas y acciones ve cada perfil. Catálogos, Avisos y Configuración respetan ese permiso (UI y API). RESERVADO no se asigna por perfil.",
        "Control de acceso por perfil.",
        "Aplicado",
    ),
    (
        "Módulo RESERVADO (confidencial).",
        "Nuevo módulo solo para Administrador (root): Video, Grabaciones y Conversaciones en solo lectura. Icono de escudo; subtítulo Confidencial. Admins de región/zona/unidad no lo ven.",
        "Área confidencial restringida a root.",
        "Aplicado",
    ),
    (
        "Grabaciones en RESERVADO (fuera de Config).",
        "Las grabaciones PTT de la app se guardan y consultan en RESERVADO. Se quitó la pestaña Grabaciones de Configuración (la URL antigua redirige).",
        "Reubicación de Grabaciones.",
        "Aplicado",
    ),
    (
        "Grupos: panel lateral de edición.",
        "Clic en un canal abre un panel a la derecha (foto, estado, miembros, roles). Se quitó la columna Acciones de la tabla; Eliminar permanente queda en el pie del panel (solo root). Flechas ↑/↓ recorren canales; clic fuera o Escape cierra.",
        "Rediseño de edición de canales.",
        "Aplicado",
    ),
    (
        "Grupos/Usuarios: filtros y orden de columnas.",
        "Filtros multi en encabezado (estilo Parque Vehicular), orden ▲/▼ al clic, arrastre de columnas, filtro Alcance por Región/Zona/Unidad. Al ocultar Filtros se limpia el filtrado. Toolbars fijas al scrollear.",
        "Mejora de listados administrativos.",
        "Aplicado",
    ),
    (
        "Solo escucha: video con imagen, sin mic.",
        "El rol «Solo escucha» puede entrar a videollamada de grupo y transmitir cámara; el sistema no le da micrófono. El PTT de radio sigue sin publicar audio.",
        "Video permitido; audio de radio bloqueado.",
        "Aplicado",
    ),
    (
        "Chats en despacho como página.",
        "Chats deja de ser solo un panel flotante: hay entrada de menú y página dedicada en la consola de despacho.",
        "Navegación de chats en despacho.",
        "Aplicado",
    ),
    (
        "Login tema obscuro: logo y transparencia.",
        "Se ajustó el arte del login obscuro (logo SICOM, esquina, transparencias) para que coincida con la referencia institucional.",
        "Pulido visual del login obscuro.",
        "Aplicado",
    ),
    (
        "Radio web: layout clásico conservado.",
        "Tras pruebas de barra de voz y controles arriba, Radio vuelve al layout habitual (cuatro columnas + PTT). La barra de voz queda guardada en código por si se pide más adelante.",
        "Restauración del módulo Radio.",
        "Aplicado",
    ),
]


def set_cell_text(cell, text, *, bold=False):
    # Keep first paragraph; clear runs
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    # remove extra paragraphs
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    # remove existing shd
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    tcPr.append(shd)


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)

    # Header meta table
    t0 = doc.tables[0]
    set_cell_text(
        t0.rows[1].cells[0],
        "“Sistema de comunicación instantánea”, Radio, Chat, Videollamada, GPS, Botón de alertas, envío de archivos y mapeo.",
    )
    set_cell_text(
        t0.rows[1].cells[1],
        "Cap. 1/o. I.C.I. Jorge Ignacio Luna García.",
    )
    set_cell_text(t0.rows[1].cells[2], FECHA)

    t1 = doc.tables[1]
    # Remove data rows (keep header)
    tbl = t1._tbl
    for tr in list(tbl.tr_lst)[1:]:
        tbl.remove(tr)

    header_tr = tbl.tr_lst[0]
    for i, (tipo, desc, accion, estado) in enumerate(CAMBIOS, start=1):
        new_tr = deepcopy(header_tr)
        tbl.append(new_tr)
        # bind as row via table rows
        row = t1.rows[-1]
        values = [str(i), tipo, DEV, desc, accion, estado]
        for cell, val in zip(row.cells, values):
            set_cell_text(cell, val, bold=(cell is row.cells[0]))
            shade_cell(cell, "F3FBF9")

    os.makedirs(os.path.dirname(OUT2), exist_ok=True)
    doc.save(OUT)
    shutil.copy2(OUT, OUT2)
    print("OK", OUT)
    print("OK", OUT2)
    print("filas", len(CAMBIOS))


if __name__ == "__main__":
    main()
