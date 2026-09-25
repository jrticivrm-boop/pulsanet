# -*- coding: utf-8 -*-
"""Alcance 5 — estado vigente: usuarios, grupos y mapa (TacticalPtx)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = r"C:\Users\JfaRgnl_TIC\Desktop\Alcance 5.docx"
OUT2 = r"C:\pulsanet_soporte\Documentos\Alcance 5.docx"
OUT_MD = r"C:\pulsanet\docs\ALCANCE_5.md"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2)
    s.right_margin = Cm(2)

GREEN = RGBColor(0x1B, 0x7A, 0x3A)
BLUE = RGBColor(0x0B, 0x3D, 0x91)
GRAY = RGBColor(0x55, 0x55, 0x55)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
BLACK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, *, bold=False, underline=False, size=11, color=None, italic=False):
    run.bold = bold
    run.underline = underline
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def shade_p(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def title(t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run(r, bold=True, size=18, color=BLUE, underline=True)


def sub(t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run(r, bold=True, size=12, color=TEAL, italic=True)


def h1(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, bold=True, size=14, color=BLUE, underline=True)
    p.paragraph_format.space_before = Pt(14)


def h2(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, bold=True, size=12, color=BLUE)
    p.paragraph_format.space_before = Pt(10)


def body(text, **kw):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(
        r,
        size=11,
        color=kw.get("color", BLACK),
        bold=kw.get("bold", False),
        italic=kw.get("italic", False),
    )


def ok(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade_p(p, "E8F5E9")
    r0 = p.add_run("VIGENTE: ")
    set_run(r0, bold=True, size=10, color=GREEN)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=GREEN)


def bullet(label, text):
    p = doc.add_paragraph(style="List Bullet")
    r1 = p.add_run(label)
    set_run(r1, bold=True, underline=True, size=11, color=BLUE)
    r2 = p.add_run(" " + text)
    set_run(r2, size=11, color=BLACK)


def add_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        rp = c.paragraphs[0]
        rr = rp.add_run(h)
        set_run(rr, bold=True, size=9, color=WHITE)
        shade_cell(c, "0D6E6E")
    for vals in rows:
        row = t.add_row().cells
        for i, val in enumerate(vals):
            row[i].text = ""
            rp = row[i].paragraphs[0]
            rr = rp.add_run(val)
            set_run(rr, bold=(i == 0), size=8, color=BLACK)
    return t


# ── Contenido ──────────────────────────────────────────────
title("Alcance 5 — Estado vigente del sistema")
sub("Registro de usuarios · Creación de grupos · Visualización en mapa")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("24-sep-2026 · Referencia de lo ya definido e implementado en TacticalPtx / SICOM")
set_run(r, size=9, color=GRAY, italic=True)

p = doc.add_paragraph()
shade_p(p, "E8F0FE")
r = p.add_run(
    "Este documento consolida las reglas actuales (Alcances 3–4 + ajustes de grupos/membresía). "
    "Sustituye a Alcance 4 como referencia de comportamiento en pantalla. "
    "No sustituye el inventario de datos «ajustes manuales» si aún hay registros por corregir en BD."
)
set_run(r, size=11, color=BLACK)

h1("1. Tres capas (no confundir)")
bullet(
    "Pertenencia:",
    "Dónde queda «anclado» el usuario al alta (región / zona / unidad). Define territorio y filtros.",
)
bullet(
    "Mapa:",
    "A quién ves en el mapa / seguimiento: matriz de roles + territorio GPS. Nadie oculta ubicación.",
)
bullet(
    "Canal (grupo / radio):",
    "Canal PTT con alcance región|zona|unidad. Hablar exige ser miembro. Quien agrega no puede «subir» de nivel.",
)
ok("Pertenencia ≠ mapa ≠ canal. Estar en el mapa de una región no abre solo los canales de radio.")

h1("2. Registro / alta de usuarios (pertenencia)")

h2("2.1 Quién puede crear a quién")
body(
    "Al asignar perfil/rol, solo se ofrecen opciones «hacia abajo» (el Administrador ve todos):"
)
add_table(
    ["Quién da de alta", "Puede crear"],
    [
        ("Administrador (root)", "Cualquier perfil"),
        ("Admin de región", "Usuario región, admin/usuario zona, admin/usuario unidad"),
        ("Admin de zona", "Usuario zona, admin/usuario unidad"),
        ("Admin de unidad", "Solo usuario de unidad"),
    ],
)

h2("2.2 Qué elige en el formulario (adscripción)")
add_table(
    ["Perfil del nuevo usuario", "Selectores obligatorios", "Cómo se guarda"],
    [
        ("Administrador", "— (global)", "Sin unit_id ni admin_scope"),
        (
            "Admin / Usuario de región",
            "Solo Región",
            "Admin → admin_scope = región; Usuario → unit_id = región",
        ),
        (
            "Admin / Usuario de zona",
            "Región + Zona",
            "Admin → admin_scope = zona; Usuario → unit_id = zona",
        ),
        (
            "Admin / Usuario de unidad",
            "Región + Zona + Unidad",
            "Admin → unit_id y admin_scope = unidad; Usuario → unit_id = unidad",
        ),
    ],
)
ok("Admin/usuario de región: no se fuerza zona ni unidad en el alta (zonas/unidades se usan al crear canales).")
ok("No existe «ocultar ubicación»: la API rechaza location-share; la UI no ofrece ese menú.")

h1("3. Creación de grupos / canales")

h2("3.1 Alcance del canal según quien lo crea")
add_table(
    ["Creador", "Alcances posibles"],
    [
        (
            "Administrador / Admin región",
            "Región (solo región) · Zona (región+zona) · Unidad (región+zona+unidad). Zona y unidad opcionales.",
        ),
        (
            "Admin de zona",
            "Canal de toda su zona (unidad vacía) o de una unidad de su zona.",
        ),
        ("Admin de unidad", "Solo canal de su unidad (fijo)."),
        ("Usuarios (región/zona/unidad)", "No crean canales de consola."),
    ],
)

h2("3.2 Quién puede agregar a quién (jerarquía)")
body("Un admin no mete perfiles «más arriba» que él:")
add_table(
    ["Quién asigna miembros", "Puede agregar"],
    [
        ("Administrador / Admin región", "Casi cualquier perfil (sujeto a geo del canal)"),
        ("Admin de zona", "Admin/usuario zona y unidad (no región)"),
        ("Admin de unidad", "Solo admin/usuario de unidad"),
    ],
)
ok("Ejemplo: admin de unidad NO puede agregar a un admin de zona.")

h2("3.3 Quién encaja en el canal (territorio + rol)")
bullet(
    "Canal de región:",
    "Personas cuya adscripción cae bajo esa región; perfiles de región siempre elegibles (quien asigna filtra aparte).",
)
bullet(
    "Canal de zona:",
    "Adscripción bajo esa zona; perfiles de región también pueden entrar si quien asigna lo permite.",
)
bullet(
    "Canal de unidad:",
    "Misma unidad; o administrador de zona cuya zona contiene esa unidad. "
    "Usuario de zona NO entra a canal de unidad (solo por ser de la zona).",
)
ok("El selector «Agregar persona…» oculta a quienes ya son miembros del grupo.")
ok("UI y API alineadas: usuario de zona no aparece para canales de unidad.")

h2("3.4 Radio PTT")
body(
    "Para hablar / escuchar en un canal hay que ser miembro. "
    "Ver gente en el mapa de la región no otorga PTT en todos los canales."
)

h1("4. Visualización en el mapa")

h2("4.1 Matriz de roles (a quién puedes ver por tipo)")
body(
    "Primero se aplica la matriz (abajo no ve arriba). Luego el territorio GPS recorta a tu región/zona/unidad."
)
add_table(
    ["Tú eres…", "En el mapa puedes ver (por rol)"],
    [
        ("Administrador", "Todos"),
        ("Admin de región", "Admin/usuario región, zona y unidad (su territorio)"),
        ("Usuario de región", "Igual: región ↓ zona ↓ unidad de su región"),
        ("Admin de zona", "Zona y unidad (su zona)"),
        ("Usuario de zona", "Solo otros usuarios de zona (pares)"),
        ("Admin de unidad", "Admin y usuarios de su unidad"),
        ("Usuario de unidad", "Solo otros usuarios de unidad (pares)"),
    ],
)

h2("4.2 Territorio GPS (dónde miras)")
bullet("Root / admin región (sin ancla rota):", "Puede ser org amplia o acotada a su región (admin_scope).")
bullet("Usuario de región:", "Solo el árbol de su región (unit_id), no toda la organización a ciegas.")
bullet("Admin / usuario de zona:", "Unidades bajo su zona.")
bullet("Admin / usuario de unidad:", "Su unidad (y subordinadas reales si aplica).")
ok("Te ves siempre a ti mismo. Nadie «oculta» el punto en el mapa por configuración de privacidad.")

h1("5. Tabla resumen (todo junto)")
add_table(
    ["Rol", "Alta (pertenencia)", "Mapa", "Crea canales", "Agrega miembros"],
    [
        ("Administrador", "Global", "Todos", "Cualquier alcance", "Cualquiera (con geo)"),
        ("Admin región", "Solo región", "Su región ↓", "R / Z / U", "Hasta su nivel"),
        ("Usuario región", "Solo región", "Su región ↓", "No", "—"),
        ("Admin zona", "R + zona", "Su zona ↓", "Zona o unidad", "Zona y unidad"),
        ("Usuario zona", "R + zona", "Pares zona", "No", "—"),
        ("Admin unidad", "R+Z+unidad", "Su unidad", "Solo su unidad", "Solo unidad"),
        ("Usuario unidad", "R+Z+unidad", "Pares unidad", "No", "—"),
    ],
)

h1("6. Cómo validar en pantalla")
bullet("1.", "Alta admin/usuario de región: solo selector Región.")
bullet("2.", "Admin unidad abre un canal de su unidad: en «Agregar persona» no debe salir un admin de zona ajeno ni un usuario de zona.")
bullet("3.", "Root/admin región en canal de unidad: sí puede salir un admin de zona de la zona padre.")
bullet("4.", "Mapa con usuario de unidad: no ve admins de zona/región.")
bullet("5.", "Mapa con usuario de región: ve marcadores de zona/unidad de su región.")
bullet("6.", "Tras agregar a alguien al canal, deja de aparecer en el combo «Agregar persona…».")

h1("7. Referencia técnica (código)")
body(
    "Visibilidad mapa: backend/src/services/visibility.js (DEFAULT_SEE, canSeePerson). "
    "Territorio track: orgUnits.js → loadTrackScope. "
    "Grupos: groupPolicy.js (memberFitsGroupGeo, actorCanAssignMember) + DispatchGroups.jsx. "
    "Alta usuarios: DispatchUsers.jsx → orgScopeForSave / allowedRoleOptions."
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— Fin · Alcance 5 —")
set_run(r, italic=True, size=9, color=GRAY)

os.makedirs(r"C:\pulsanet_soporte\Documentos", exist_ok=True)
for path in (OUT, OUT2):
    try:
        doc.save(path)
        print("OK", path)
    except PermissionError:
        alt = path.replace(".docx", "_nuevo.docx")
        doc.save(alt)
        print("SKIP abierto →", alt)

MD = """# Alcance 5 — Estado vigente del sistema

**24-sep-2026** · Registro de usuarios · Creación de grupos · Visualización en mapa  
Referencia de lo ya definido e implementado en TacticalPtx / SICOM.

Consolida Alcances 3–4 y los ajustes recientes de membresía en grupos.

## 1. Tres capas (no confundir)

| Capa | Qué es |
|------|--------|
| **Pertenencia** | Dónde queda anclado el usuario al alta (región / zona / unidad). |
| **Mapa** | A quién ves: matriz de roles + territorio GPS. Nadie oculta ubicación. |
| **Canal** | Grupo/radio PTT con alcance región\|zona\|unidad. Hablar = ser miembro. |

Pertenencia ≠ mapa ≠ canal.

## 2. Registro / alta de usuarios

### Quién puede crear a quién

| Quién da de alta | Puede crear |
|------------------|---------------|
| Administrador (root) | Cualquier perfil |
| Admin de región | Usuario región, admin/usuario zona, admin/usuario unidad |
| Admin de zona | Usuario zona, admin/usuario unidad |
| Admin de unidad | Solo usuario de unidad |

### Adscripción al guardar

| Perfil | Selectores | Guardado |
|--------|------------|----------|
| Administrador | — | Sin unit / scope |
| Admin / Usuario región | Solo región | Admin: `admin_scope`=región · Usuario: `unit_id`=región |
| Admin / Usuario zona | Región + zona | Admin: `admin_scope`=zona · Usuario: `unit_id`=zona |
| Admin / Usuario unidad | Región + zona + unidad | Admin: ambos = unidad · Usuario: `unit_id`=unidad |

- Región: no se fuerza zona/unidad en el alta.
- No hay «ocultar ubicación».

## 3. Creación de grupos / canales

### Alcance según creador

| Creador | Alcances |
|---------|----------|
| Root / admin región | Región · Zona · Unidad (zona/unidad opcionales) |
| Admin zona | Toda su zona o una unidad de su zona |
| Admin unidad | Solo su unidad |

### Quién agrega a quién

| Quién asigna | Puede agregar |
|--------------|---------------|
| Root / admin región | Casi cualquiera (con filtro geo) |
| Admin zona | Zona y unidad (no región) |
| Admin unidad | Solo admin/usuario de unidad |

**Ejemplo:** admin de unidad **no** puede agregar a un admin de zona.

### Encaje en el canal

- **Región / zona:** adscripción bajo el ancla; perfiles de región pueden entrar si quien asigna lo permite.
- **Unidad:** misma unidad, o **admin de zona** cuya zona contiene esa unidad. **Usuario de zona no** entra solo por ser de la zona.
- El combo «Agregar persona…» **oculta** a quien ya es miembro.
- Hablar en radio exige membresía (el mapa no abre PTT solo).

## 4. Mapa — a quién ves

### Matriz de roles

| Tú eres… | Ves (por rol) |
|----------|----------------|
| Administrador | Todos |
| Admin / usuario región | Región ↓ zona ↓ unidad (su región) |
| Admin zona | Zona ↓ unidad |
| Usuario zona | Solo pares usuario de zona |
| Admin unidad | Admin y usuarios de su unidad |
| Usuario unidad | Solo pares usuario de unidad |

### Territorio GPS

- Usuario de región: árbol de **su** región (no org entera a ciegas).
- Zona: unidades bajo su zona.
- Unidad: su unidad.
- Siempre te ves a ti mismo.

## 5. Validación rápida

1. Alta región → solo selector Región.  
2. Admin unidad + canal unidad → no sale admin zona ajeno ni usuario de zona.  
3. Root en canal unidad → sí puede salir admin de zona padre.  
4. Usuario unidad en mapa → no ve zona/región.  
5. Tras agregar a alguien → desaparece del combo Agregar.

## 6. Código de referencia

- `visibility.js` — matriz mapa  
- `orgUnits.js` → `loadTrackScope` — territorio  
- `groupPolicy.js` + `DispatchGroups.jsx` — canales  
- `DispatchUsers.jsx` → `orgScopeForSave` / `allowedRoleOptions` — alta  

— Fin · Alcance 5 —
"""

with open(OUT_MD, "w", encoding="utf-8") as f:
    f.write(MD)
print("OK", OUT_MD)
