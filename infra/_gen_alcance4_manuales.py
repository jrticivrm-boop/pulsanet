# -*- coding: utf-8 -*-
"""Alcance 4 — ajustes manuales.docx: solo pendientes reales (perfil vacío / org incompleta)."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
SOPORTE = Path(r"C:\pulsanet_soporte\Documentos")
DATA = Path(r"C:\pulsanet\infra\_alcance4_manual_data.json")
DOCX_NAME = "Alcance 4 — ajustes manuales.docx"

BLUE = RGBColor(0x0B, 0x3D, 0x91)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
GRAY = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x22, 0x22, 0x22)
ORANGE = RGBColor(0xB8, 0x5C, 0x00)
GREEN = RGBColor(0x1B, 0x7A, 0x3A)
RED = RGBColor(0xB0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, *, bold=False, underline=False, size=11, color=None, italic=False):
    run.bold = bold
    run.underline = underline
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def shade_p(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def cell_text(cell, text, *, bold=False, size=9, color=BLACK):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_run(r, bold=bold, size=size, color=color)


def title(doc, t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run(r, bold=True, size=18, color=BLUE, underline=True)


def sub(doc, t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run(r, bold=True, size=12, color=TEAL, italic=True)


def h1(doc, t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, bold=True, size=14, color=BLUE, underline=True)
    p.paragraph_format.space_before = Pt(14)


def h2(doc, t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, bold=True, size=12, color=BLUE)
    p.paragraph_format.space_before = Pt(10)


def body(doc, text, *, color=BLACK, bold=False, italic=False, size=11, shade=None):
    p = doc.add_paragraph()
    if shade:
        shade_p(p, shade)
    r = p.add_run(text)
    set_run(r, size=size, color=color, bold=bold, italic=italic)


def main():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.6)
        s.right_margin = Cm(1.6)

    title(doc, "Alcance 4 — ajustes manuales")
    sub(doc, "Solo lo que SÍ debes corregir (perfil vacío o región/zona/unidad faltante)")
    body(
        doc,
        f"Generado: {data.get('generatedAt', '')} · Datos vivos de la base",
        size=9,
        color=GRAY,
        italic=True,
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    body(
        doc,
        "REGLA SIMPLE",
        bold=True,
        size=12,
        color=ORANGE,
        shade="FFF8E1",
    )
    body(
        doc,
        "1) Si en la ficha ya ves Región / Zona / Unidad correctas → NO las cambies.",
        size=11,
    )
    body(
        doc,
        "2) Si el campo «Perfil *» está vacío en la BASE (aunque el combo muestre un nombre), "
        "elige el perfil sugerido y pulsa Guardar cambios.",
        size=11,
        bold=True,
    )
    body(
        doc,
        "3) Solo si falta Región o Zona (o Unidad) según el rol → completa Alcance y Guardar.",
        size=11,
    )
    body(
        doc,
        "Nota: en Administrador de zona es NORMAL no elegir unidad. No es un error.",
        size=10,
        color=GRAY,
        italic=True,
    )

    only_p = data.get("onlyProfile") or []
    only_o = data.get("onlyOrg") or []
    both = data.get("both") or []
    groups = data.get("groups") or []

    h1(doc, "Resumen")
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    rows_sum = [
        ("Usuarios solo con PERFIL vacío (alcance OK)", str(len(only_p))),
        ("Usuarios con región/zona/unidad FALTANTE", str(len(only_o) + len(both))),
        ("Usuarios con ambos problemas", str(len(both))),
        ("Canales/grupos sin ancla", str(len(groups))),
    ]
    for i, (a, b) in enumerate(rows_sum):
        cell_text(t.rows[i].cells[0], a, bold=True, size=9, color=BLUE)
        shade_cell(t.rows[i].cells[0], "E8F0FE")
        cell_text(t.rows[i].cells[1], b, bold=True, size=11, color=RED if b != "0" else GREEN)

    # —— A ——
    h1(doc, "A) Usuarios: solo falta guardar el Perfil")
    if not only_p:
        body(doc, "Ninguno. Nada que hacer en esta sección.", color=GREEN, bold=True, shade="E8F5E9")
    else:
        body(
            doc,
            "Cómo está: Rol y Alcance (región/zona/unidad) YA están bien. "
            "Cómo debe: mismo alcance + Perfil rellenado. "
            "Dónde: Administración → Usuarios → Editar → Perfil * → elegir → Guardar cambios.",
            size=10,
            shade="E8F0FE",
        )
        tbl = doc.add_table(rows=1 + len(only_p), cols=5)
        tbl.style = "Table Grid"
        headers = ["Usuario", "Rol", "Alcance HOY (OK)", "Perfil HOY", "Perfil DEBE ser"]
        for i, h in enumerate(headers):
            cell_text(tbl.rows[0].cells[i], h, bold=True, size=8, color=WHITE)
            shade_cell(tbl.rows[0].cells[i], "0D6E6E")
        for ri, u in enumerate(only_p, start=1):
            cell_text(tbl.rows[ri].cells[0], f"{u['displayName']}\n({u['username']})", size=8)
            cell_text(tbl.rows[ri].cells[1], u["roleLabel"], size=8)
            cell_text(tbl.rows[ri].cells[2], u["orgNow"], size=8, color=GREEN)
            cell_text(tbl.rows[ri].cells[3], u["profileNow"], size=8, color=RED, bold=True)
            cell_text(tbl.rows[ri].cells[4], f"«{u['profileMust']}»", size=8, color=GREEN, bold=True)
            shade_cell(tbl.rows[ri].cells[3], "FFEBEE")
            shade_cell(tbl.rows[ri].cells[4], "E8F5E9")

        body(
            doc,
            "Ejemplo Pérez Informática (mperezh3): si el combo ya muestra «Administrador de zona» "
            "pero sigue en esta lista, es porque NO se guardó en base. Abre → confirma Perfil → Guardar cambios.",
            size=9,
            color=ORANGE,
            italic=True,
        )

    # —— B ——
    h1(doc, "B) Usuarios: falta Región / Zona / Unidad según el rol")
    org_list = only_o + both
    if not org_list:
        body(
            doc,
            "Ninguno. Todos los activos tienen región/zona/unidad coherente con su rol.",
            color=GREEN,
            bold=True,
            shade="E8F5E9",
        )
    else:
        body(
            doc,
            "Estos SÍ necesitan completar el bloque Alcance (Región → Zona → Unidad según el perfil/rol).",
            size=10,
            shade="FFF8E1",
        )
        tbl = doc.add_table(rows=1 + len(org_list), cols=5)
        tbl.style = "Table Grid"
        headers = ["Usuario", "Rol", "Cómo está el alcance", "Qué falta", "Perfil"]
        for i, h in enumerate(headers):
            cell_text(tbl.rows[0].cells[i], h, bold=True, size=8, color=WHITE)
            shade_cell(tbl.rows[0].cells[i], "B85C00")
        for ri, u in enumerate(org_list, start=1):
            cell_text(tbl.rows[ri].cells[0], f"{u['displayName']}\n({u['username']})", size=8)
            cell_text(tbl.rows[ri].cells[1], u["roleLabel"], size=8)
            cell_text(tbl.rows[ri].cells[2], u["orgNow"], size=8, color=RED)
            cell_text(tbl.rows[ri].cells[3], " · ".join(u["problems"]), size=8, color=RED, bold=True)
            cell_text(
                tbl.rows[ri].cells[4],
                u["profileNow"] if not u["needProfile"] else f"También vacío → «{u['profileMust']}»",
                size=8,
            )

    # —— C ——
    h1(doc, "C) Canales / grupos sin ancla")
    if not groups:
        body(doc, "Ninguno pendiente.", color=GREEN, bold=True)
    else:
        body(
            doc,
            "Dónde: Administración → Grupos → editar canal → fijar Región/Zona/Unidad del alcance → Guardar.",
            size=10,
        )
        tbl = doc.add_table(rows=1 + len(groups), cols=3)
        tbl.style = "Table Grid"
        for i, h in enumerate(["Canal", "CÓMO ESTÁ", "CÓMO DEBE"]):
            cell_text(tbl.rows[0].cells[i], h, bold=True, size=8, color=WHITE)
            shade_cell(tbl.rows[0].cells[i], "0D6E6E")
        for ri, g in enumerate(groups, start=1):
            cell_text(tbl.rows[ri].cells[0], g["name"], size=9, bold=True)
            cell_text(tbl.rows[ri].cells[1], g["howItIs"], size=8, color=RED)
            cell_text(tbl.rows[ri].cells[2], g["howItShouldBe"], size=8, color=GREEN)
            shade_cell(tbl.rows[ri].cells[1], "FFF3E0")
            shade_cell(tbl.rows[ri].cells[2], "E8F5E9")

    h1(doc, "D) Qué NO tocar")
    body(doc, "• No cambies región/zona/unidad de la sección A: ya están correctas.", size=10)
    body(doc, "• No busques un campo «Pertenencia» aparte: en la UI se llama Alcance (Región / Zona).", size=10)
    body(doc, "• Usuarios que no están en A ni B: no hace falta abrirlos.", size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— Fin · Alcance 4 ajustes manuales —")
    set_run(r, italic=True, size=9, color=GRAY)

    SOPORTE.mkdir(parents=True, exist_ok=True)
    for path in (DESKTOP / DOCX_NAME, SOPORTE / DOCX_NAME):
        try:
            doc.save(str(path))
            print("OK", path)
        except PermissionError:
            alt = path.with_name(path.stem + " (nuevo).docx")
            doc.save(str(alt))
            print("SKIP abierto →", alt)


if __name__ == "__main__":
    main()
