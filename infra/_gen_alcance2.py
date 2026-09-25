# -*- coding: utf-8 -*-
"""Genera Alcance 2.docx en Escritorio + pulsanet_soporte/Documentos."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = r"C:\Users\JfaRgnl_TIC\Desktop\Alcance 2.docx"
OUT2 = r"C:\pulsanet_soporte\Documentos\Alcance 2.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

GREEN = RGBColor(0x1B, 0x7A, 0x3A)
RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x0B, 0x3D, 0x91)
GRAY = RGBColor(0x55, 0x55, 0x55)
ORANGE = RGBColor(0xB8, 0x5C, 0x00)
BLACK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, *, bold=False, underline=False, size=11, color=None, italic=False):
    run.bold = bold
    run.underline = underline
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(paragraph, hex_color):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, bold=True, size=18, color=BLUE, underline=True)


def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, bold=True, size=12, color=GRAY, italic=True)


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=True, size=14, color=BLUE, underline=True)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=True, size=12, color=BLUE)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def body_parts(parts):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    for text, kw in parts:
        r = p.add_run(text)
        set_run(r, size=11, color=kw.get("color", BLACK), **{k: v for k, v in kw.items() if k != "color"})


def bullet(label, text, *, label_color=BLUE):
    p = doc.add_paragraph(style="List Bullet")
    r1 = p.add_run(label)
    set_run(r1, bold=True, underline=True, size=11, color=label_color)
    r2 = p.add_run(" " + text)
    set_run(r2, size=11, color=BLACK)


def ok(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)
    r0 = p.add_run("BIEN: ")
    set_run(r0, bold=True, size=10, color=GREEN)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=GREEN)


def gap(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)
    shade_paragraph(p, "FFE6E6")
    r0 = p.add_run("FALTA / REVISAR: ")
    set_run(r0, bold=True, underline=True, size=10, color=RED)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=RED, bold=True)


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r0 = p.add_run("Nota: ")
    set_run(r0, bold=True, italic=True, size=10, color=ORANGE)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=ORANGE, italic=True)


# ── Contenido ──
add_title("Alcance 2 — Guía clara de roles")
add_subtitle("TacticalPtx · Visibilidad en mapa · Creación de grupos")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Versión revisada del documento «Alcance» (escritorio) · 23-sep-2026")
set_run(r, size=9, color=GRAY, italic=True)

p = doc.add_paragraph()
shade_paragraph(p, "E8F0FE")
r = p.add_run("Cómo leer este documento: ")
set_run(r, bold=True, size=11, color=BLUE)
r = p.add_run("cada rol tiene tres cosas distintas. ")
set_run(r, size=11, color=BLACK)
r = p.add_run("Alcance")
set_run(r, bold=True, underline=True, size=11, color=BLUE)
r = p.add_run(" = dónde “pertenece” al darlo de alta. ")
set_run(r, size=11, color=BLACK)
r = p.add_run("Mapa")
set_run(r, bold=True, underline=True, size=11, color=BLUE)
r = p.add_run(" = a quién ve en el mapa. ")
set_run(r, size=11, color=BLACK)
r = p.add_run("Grupos")
set_run(r, bold=True, underline=True, size=11, color=BLUE)
r = p.add_run(" = si puede crear canales y con qué gente. ")
set_run(r, size=11, color=BLACK)
r = p.add_run("Verde = ya está en el sistema. Rojo = falta o no coincide.")
set_run(r, bold=True, size=11, color=RED)

h1("1. Ideas clave (para cualquier usuario)")
body_parts(
    [
        ("No es lo mismo ", {}),
        ("rol", {"bold": True, "underline": True}),
        (", ", {}),
        ("adscripción / alcance", {"bold": True, "underline": True}),
        (" y ", {}),
        ("pertenecer a un canal (grupo)", {"bold": True, "underline": True}),
        (
            ". El rol dice el poder. El alcance dice en qué región/zona/unidad vive. "
            "El grupo es un canal de radio aparte: hay que agregarlo como miembro para hablar.",
            {},
        ),
    ]
)
body_parts(
    [
        ("Jerarquía de ejemplo: ", {"bold": True}),
        ("Región IV R.M. → Zona 48/a. Z.M. → Unidad / organismo.", {}),
    ]
)

h1("2. Por cada rol")

h2("2.1 Administrador (sistema / root)")
bullet("Alcance:", "Todas las regiones, zonas y unidades. En el alta no elige cascada (queda global).")
bullet("Mapa:", "Ve a todos los usuarios.")
bullet("Grupos:", "Puede crear canales de cualquier alcance y agregar usuarios disponibles.")
ok("Coincide con el sistema actual.")

h2("2.2 Administrador de región")
bullet(
    "Alcance (según documento original):",
    "Solo la región (ej. IV R.M.) y ahí termina la selección.",
)
bullet("Mapa:", "Ve a todos los usuarios de esa región (zonas y unidades hijas).")
bullet(
    "Grupos:",
    "Puede crear canales de su región. Si elige «Todas las zonas», puede agregar gente de todas "
    "las unidades de la región. Si elige una zona concreta, debe bajar a una unidad y solo "
    "agregar gente de esa unidad.",
)
ok("Puede administrar su región y crear canales de región / zona / unidad.")
gap(
    "En el alta actual el admin de región aún elige «todas las zonas» o «una zona». "
    "El documento pide que la selección termine solo en la región."
)
gap(
    "Al crear un grupo: si ya eligió una zona concreta, el sistema todavía permite "
    "«todos los organismos de la zona». El documento pide obligar a elegir una unidad "
    "y solo miembros de esa unidad."
)

h2("2.3 Usuario de región")
bullet("Alcance:", "La región seleccionada (ej. IV R.M.).")
bullet("Mapa (documento):", "Debe ver a todos los usuarios de esa región (zonas y unidades).")
bullet("Grupos:", "No puede crear grupos.")
ok("No puede crear canales.")
gap(
    "Mapa: hoy el usuario de región solo ve, por matriz, a otros de nivel región "
    "(admins/usuarios de región). NO ve por defecto a gente de zona ni de unidad. "
    "Eso NO cumple el documento."
)
gap(
    "GPS / seguimiento: el sistema trata al usuario de región como alcance amplio (orgWide) "
    "y no respeta del todo la región elegida en la cascada del alta."
)
gap(
    "Radio: aunque sea «de región», solo oye canales donde lo metieron como miembro. "
    "El alcance de región no le abre la radio solo por jerarquía."
)

h2("2.4 Administrador de zona")
bullet("Alcance:", "Región + una zona concreta. Hay que quitar «Todas las zonas». Ahí termina.")
bullet("Mapa:", "Ve a todos los de su zona y de las unidades de esa zona.")
bullet(
    "Grupos:",
    "Región y zona bloqueadas a las suyas. Elige una unidad o «Todas las unidades» "
    "para armar el canal y agregar miembros.",
)
ok("En el alta de este rol ya no hay «todas las zonas»: debe elegir una zona.")
ok("Creación de grupos: región/zona fijadas; puede canal de toda la zona o de una unidad.")
note("Comprobar en pantalla que al crear el canal se vean bloqueados región y zona con los nombres correctos.")

h2("2.5 Administrador de unidad")
bullet("Alcance:", "Región → Zona → Unidad. Ahí termina.")
bullet("Mapa:", "Solo usuarios de su unidad.")
bullet("Grupos:", "Todo bloqueado a su unidad. Solo nombra el canal y agrega gente de su unidad.")
ok("Coincide con el sistema actual.")

h2("2.6 Usuario de unidad")
bullet("Alcance:", "Región → Zona → Unidad.")
bullet("Mapa:", "Solo su unidad.")
bullet("Grupos:", "No puede crear grupos.")
ok("Coincide con el sistema actual.")

h2("2.7 Usuario de zona (no estaba en el documento original)")
body_parts(
    [
        ("El documento original ", {}),
        ("no menciona", {"bold": True, "underline": True}),
        (" al ", {}),
        ("Usuario de zona", {"bold": True}),
        (". En el sistema sí existe (entre admin de zona y usuario de unidad).", {}),
    ]
)
gap(
    "Falta definir en el documento oficial: alcance (región+zona), mapa "
    "(¿solo usuarios de zona?), y que no crea grupos. Hoy el código lo trata como par de zona."
)

h1("3. Ocultar ubicación (admins)")
body_parts(
    [
        ("Quién puede ocultar: ", {"bold": True, "underline": True}),
        ("Administrador, Admin de región, Admin de zona y Admin de unidad.", {}),
    ]
)
body_parts(
    [
        ("Quién NO puede ocultar: ", {"bold": True, "underline": True}),
        ("Usuario de región, Usuario de zona y Usuario de unidad.", {}),
    ]
)
body_parts(
    [
        ("Regla importante: ", {"bold": True}),
        (
            "aunque un admin oculte su ubicación, un superior jerárquico sigue viéndolo. "
            "Ejemplo: admin de zona oculto → invisible para sus unidades, pero visible "
            "para admin de región.",
            {},
        ),
    ]
)
ok("Los usuarios (*_user) no pueden ocultar ubicación.")
ok("Un superior admin suele seguir viendo a un inferior oculto.")
gap(
    "El documento dice que el USUARIO DE REGIÓN también debe seguir viendo al admin de zona "
    "aunque este oculte su ubicación. En el código actual eso está bloqueado a propósito "
    "(el usuario de región NO ve al admin de zona oculto). Hay que decidir y corregir."
)

h1("4. Resumen rápido: ¿cumplimos el documento?")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
for i, h in enumerate(["Tema del documento", "¿Está?", "Comentario"]):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    r = p.add_run(h)
    set_run(r, bold=True, size=10, color=WHITE)
    shade_cell(hdr[i], "0B3D91")

rows_data = [
    ("Admin (root): todo / ve todos / crea grupos", "SÍ", "OK"),
    ("Admin región: mapa de su región", "SÍ (aprox.)", "OK funcional"),
    ("Admin región: alta solo hasta región", "NO", "Hoy pide zona o «todas»"),
    ("Admin región: grupo con zona → forzar unidad", "NO", "Permite toda la zona"),
    ("Usuario región: no crea grupos", "SÍ", "OK"),
    ("Usuario región: mapa ve zonas y unidades", "NO", "Matriz solo ve nivel región"),
    ("Usuario región: radio por jerarquía", "NO", "Solo si es miembro del canal"),
    ("Admin zona: sin «todas las zonas» en su alta", "SÍ", "OK"),
    ("Admin zona: grupos con unidad o todas", "SÍ", "OK"),
    ("Admin / usuario unidad", "SÍ", "OK"),
    ("Ocultar ubicación solo admins", "SÍ", "OK"),
    ("Usuario región ve admin zona oculto", "NO", "Código lo impide"),
    ("Usuario de zona en el documento", "NO definido", "Existe en sistema"),
]

for tema, estado, com in rows_data:
    row = table.add_row().cells
    for i, val in enumerate((tema, estado, com)):
        row[i].text = ""
        p = row[i].paragraphs[0]
        color = BLACK
        bold = False
        if i == 1:
            bold = True
            if val.startswith("SÍ"):
                color = GREEN
            elif val.startswith("NO"):
                color = RED
            else:
                color = ORANGE
        r = p.add_run(val)
        set_run(r, bold=bold, size=9, color=color)

h1("5. Qué falta hacer (lista de trabajo)")
gap(
    "1) Ajustar alta de Administrador de región: selección termina en la región "
    "(sin forzar zona), si así lo confirma el documento."
)
gap(
    "2) Ajustar creación de grupos del admin de región: si elige una zona, obligar unidad "
    "(quitar «todos los organismos» en ese caso)."
)
gap(
    "3) Hacer que el Usuario de región VEA en el mapa a zonas y unidades de su región "
    "(cambiar matriz / can_see)."
)
gap("4) Alinear GPS del usuario de región con la región elegida (hoy orgWide).")
gap(
    "5) Definir si el usuario de región debe oír canales por jerarquía o solo por membresía "
    "(hoy solo membresía)."
)
gap(
    "6) Decidir si el usuario de región ve a un admin de zona con ubicación oculta "
    "(documento dice que sí; código dice que no)."
)
gap("7) Documentar formalmente el Usuario de zona (alcance, mapa, grupos).")
note(
    "Pendientes ya conocidos del producto: visibilidad «Ver región/zonas/unidades» en el alta "
    "es decorativa (el backend la pisa con defaults del rol)."
)

h1("6. Glosario breve")
bullet("Región:", "Nivel más alto del árbol (ej. IV R.M.).")
bullet("Zona / C.G.:", "Nivel medio bajo la región (ej. 48/a. Z.M.).")
bullet("Unidad:", "Organismo o servicio desplegado.")
bullet("Canal / grupo:", "Sala de radio PTT. No es lo mismo que «pertenecer a una zona».")
bullet(
    "Perfil de acceso:",
    "Plantilla de permisos de módulos; en el alta reciente se elige perfil y de ahí sale el rol.",
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— Fin de Alcance 2 —")
set_run(r, italic=True, size=9, color=GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Basado en: Alcance.docx (escritorio) + revisión del código TacticalPtx (sept 2026)."
)
set_run(r, italic=True, size=8, color=GRAY)

os.makedirs(r"C:\pulsanet_soporte\Documentos", exist_ok=True)
doc.save(OUT)
doc.save(OUT2)
print("OK", OUT)
print("OK", OUT2)
