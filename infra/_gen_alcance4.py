# -*- coding: utf-8 -*-
"""Alcance 4 — cómo quedó implementado en TacticalPtx."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = r"C:\Users\JfaRgnl_TIC\Desktop\Alcance 4.docx"
OUT2 = r"C:\pulsanet_soporte\Documentos\Alcance 4.docx"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2)
    s.right_margin = Cm(2)

GREEN = RGBColor(0x1B, 0x7A, 0x3A)
RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x0B, 0x3D, 0x91)
GRAY = RGBColor(0x55, 0x55, 0x55)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
BLACK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xB8, 0x5C, 0x00)


def set_run(run, *, bold=False, underline=False, size=11, color=None, italic=False):
    run.bold = bold
    run.underline = underline
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def shade_p(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def title(t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run(r, bold=True, size=18, color=BLUE, underline=True)


def sub(t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run(r, bold=True, size=12, color=TEAL, italic=True)


def h1(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, bold=True, size=14, color=BLUE, underline=True)
    p.paragraph_format.space_before = Pt(14)


def h2(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, bold=True, size=12, color=BLUE)
    p.paragraph_format.space_before = Pt(10)


def body(text, **kw):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=11, color=kw.get("color", BLACK), bold=kw.get("bold", False), italic=kw.get("italic", False))


def ok(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade_p(p, "E8F5E9")
    r0 = p.add_run("IMPLEMENTADO: ")
    set_run(r0, bold=True, size=10, color=GREEN)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=GREEN)


def manual(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade_p(p, "FFF8E1")
    r0 = p.add_run("AJUSTE MANUAL: ")
    set_run(r0, bold=True, underline=True, size=10, color=ORANGE)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=ORANGE, bold=True)


def bullet(label, text):
    p = doc.add_paragraph(style="List Bullet")
    r1 = p.add_run(label)
    set_run(r1, bold=True, underline=True, size=11, color=BLUE)
    r2 = p.add_run(" " + text)
    set_run(r2, size=11, color=BLACK)


title("Alcance 4 — Cómo quedó el sistema")
sub("Reglas de pertenencia y canales · Implementación TacticalPtx")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("23-sep-2026 · Sustituye Alcance 3 como referencia de lo ya desplegado en código")
set_run(r, size=9, color=GRAY, italic=True)

p = doc.add_paragraph()
shade_p(p, "E8F0FE")
r = p.add_run(
    "Este documento describe el comportamiento real tras implementar Alcance 3. "
    "Los ajustes de datos (usuarios/grupos ya existentes) están en el archivo aparte "
    "«Alcance 4 — ajustes manuales.docx» (Word)."
)
set_run(r, size=11, color=BLACK)

h1("1. Tres capas (vigentes)")
bullet("Pertenencia:", "Dónde vive la persona al alta (región / zona / unidad).")
bullet("Mapa:", "A quién ve: solo jerarquía (abajo no ve arriba). Nadie oculta ubicación.")
bullet("Canal:", "Radio PTT; crear canal tiene cascada propia; hablar exige membresía.")
ok("Separación pertenencia ≠ mapa ≠ canal documentada en UI de Usuarios y Grupos.")

h1("2. Qué se implementó en código")
ok("Ocultar ubicación desactivado: API rechaza location-share; UI sin menú «Ocultar»; canSeePerson ignora location_share.")
ok("Alta / edición region_admin y region_user: solo eligen región (sin zona obligatoria).")
ok("Usuario de región: matriz de mapa incluye zona y unidad; GPS/track acotado a su región (unit_id), no orgWide ciego.")
ok("Grupos admin de región: misma lógica de niveles (región / zona+todas unidades / una unidad); textos de ayuda actualizados.")
ok("Jerarquía: usuario de unidad no ve zona ni región (matriz).")
ok("Usuario de zona formalizado en ayudas de rol.")

h1("3. Tabla de verdad (como quedó)")

t = doc.add_table(rows=1, cols=4)
t.style = "Table Grid"
for i, h in enumerate(["Rol", "Pertenencia (alta)", "Mapa", "Crea canales"]):
    c = t.rows[0].cells[i]
    c.text = ""
    rp = c.paragraphs[0]
    rr = rp.add_run(h)
    set_run(rr, bold=True, size=9, color=WHITE)
    shade_cell(c, "0D6E6E")

rows = [
    ("Administrador", "Global", "Todos", "Cualquier alcance"),
    ("Admin región", "Solo región", "Toda su región ↓", "Región / zona / unidad"),
    ("Usuario región", "Solo región", "Toda su región ↓", "No"),
    ("Admin zona", "Región + zona", "Su zona ↓", "Zona o unidad"),
    ("Usuario zona", "Región + zona", "Pares de zona", "No"),
    ("Admin unidad", "R+Z+unidad", "Su unidad", "Solo su unidad"),
    ("Usuario unidad", "R+Z+unidad", "Pares de unidad", "No"),
]
for vals in rows:
    row = t.add_row().cells
    for i, val in enumerate(vals):
        row[i].text = ""
        rp = row[i].paragraphs[0]
        rr = rp.add_run(val)
        set_run(rr, bold=(i == 0), size=8, color=BLACK)

h1("4. Radio (sin cambio de modelo)")
body(
    "Seguir siendo miembro del canal es obligatorio para PTT. "
    "El mapa amplio de región no abre canales solos."
)
ok("App móvil sigue listando canales con membersOnly.")

h1("5. Datos que tú debes revisar")
manual(
    "Abrir «Alcance 4 — ajustes manuales.docx» en el Escritorio. "
    "Ahí está el inventario vivo de la base: pertenencias de región incorrectas, "
    "usuarios sin profile_id, canales sin ancla, etc."
)
body(
    "Resumen al momento de generar este Word: había al menos un admin de región sin ancla de región; "
    "varios usuarios activos sin perfil; y un par de canales con alcance unit sin unidad.",
    italic=True,
    color=GRAY,
)

h1("6. Archivos tocados (referencia técnica)")
body(
    "Backend: visibility.js, roles.js, orgUnits.js (loadTrackScope), profiles.js (ruta location-share), "
    "_test_visibility.js. Frontend: DispatchUsers.jsx (cascada región, sin menú ubicación), "
    "DispatchGroups.jsx (textos), DispatchProfiles.jsx (copy)."
)

h1("7. Cómo validar en pantalla")
bullet("1.", "Alta de Administrador / Usuario de región: solo selector Región.")
bullet("2.", "Más → Usuarios: ya no aparece «Ubicación / Ocultar».")
bullet("3.", "Mapa con usuario de región: debe ver marcadores de zona/unidad de su región.")
bullet("4.", "Grupos (admin región): Todas las zonas · o zona + todas unidades · o una unidad.")
bullet("5.", "Usuario de unidad: no ve admins de zona/región en mapa.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— Fin · Alcance 4 —")
set_run(r, italic=True, size=9, color=GRAY)

os.makedirs(r"C:\pulsanet_soporte\Documentos", exist_ok=True)
for path in (OUT, OUT2):
    try:
        doc.save(path)
        print("OK", path)
    except PermissionError:
        print("SKIP (archivo abierto):", path)
