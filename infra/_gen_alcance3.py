# -*- coding: utf-8 -*-
"""Genera Alcance 3 — documento normativo corregido (sin ocultar ubicación)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = r"C:\Users\JfaRgnl_TIC\Desktop\Alcance 3.docx"
OUT2 = r"C:\pulsanet_soporte\Documentos\Alcance 3.docx"

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

GREEN = RGBColor(0x1B, 0x7A, 0x3A)
RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x0B, 0x3D, 0x91)
GRAY = RGBColor(0x55, 0x55, 0x55)
ORANGE = RGBColor(0xB8, 0x5C, 0x00)
BLACK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)


def set_run(run, *, bold=False, underline=False, size=11, color=None, italic=False):
    run.bold = bold
    run.underline = underline
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def shade_p(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, bold=True, size=18, color=BLUE, underline=True)


def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, bold=True, size=12, color=TEAL, italic=True)


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=True, size=14, color=BLUE, underline=True)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=True, size=12, color=BLUE)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def p_parts(parts, *, shade=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if shade:
        shade_p(p, shade)
    for text, kw in parts:
        r = p.add_run(text)
        set_run(
            r,
            size=kw.get("size", 11),
            color=kw.get("color", BLACK),
            bold=kw.get("bold", False),
            underline=kw.get("underline", False),
            italic=kw.get("italic", False),
        )
    return p


def bullet(label, text, *, label_color=BLUE):
    p = doc.add_paragraph(style="List Bullet")
    r1 = p.add_run(label)
    set_run(r1, bold=True, underline=True, size=11, color=label_color)
    r2 = p.add_run(" " + text)
    set_run(r2, size=11, color=BLACK)


def rule(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade_p(p, "E8F5E9")
    r0 = p.add_run("REGLA: ")
    set_run(r0, bold=True, size=10, color=GREEN)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=GREEN)


def proposal(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade_p(p, "FFF8E1")
    r0 = p.add_run("PROPUESTA (integrada): ")
    set_run(r0, bold=True, underline=True, size=10, color=ORANGE)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=ORANGE, bold=True)


def warn(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade_p(p, "FFE6E6")
    r0 = p.add_run("PENDIENTE EN CÓDIGO: ")
    set_run(r0, bold=True, underline=True, size=10, color=RED)
    r1 = p.add_run(text)
    set_run(r1, size=10, color=RED, bold=True)


# ═══════════════ DOCUMENTO ═══════════════
title("Reglas de pertenencia y canales")
subtitle("Documento oficial · Alcance 3 · TacticalPtx")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Nombre propuesto: «Reglas de pertenencia y canales»  ·  "
    "Sustituye Alcance / Alcance 2  ·  23-sep-2026"
)
set_run(r, size=9, color=GRAY, italic=True)

p_parts(
    [
        (
            "Este documento corrige malentendidos del Alcance original: "
            "separa qué es la persona, qué es el mapa y qué es un canal; "
            "elimina ocultar ubicación; y deja reglas claras para admin de región "
            "(alta vs grupos) sin contradicciones.",
            {"size": 11},
        )
    ],
    shade="E8F0FE",
)

h1("0. Nombre del documento")
p_parts(
    [
        ("Nombre corto: ", {"bold": True}),
        ("Alcance 3", {"bold": True, "color": BLUE}),
    ]
)
p_parts(
    [
        ("Nombre oficial propuesto: ", {"bold": True}),
        ("«Reglas de pertenencia y canales»", {"bold": True, "underline": True, "color": TEAL}),
    ]
)
p_parts(
    [
        (
            "Por qué: «Alcance» solo confundía. Aquí se habla de tres capas: "
            "dónde pertenece la persona, a quién ve en el mapa, y qué canales puede crear o usar.",
            {"italic": True, "color": GRAY},
        )
    ]
)

h1("1. Tres capas (léelo primero — evita confusiones)")
p_parts(
    [
        ("Capa A — Pertenencia (adscripción)", {"bold": True, "underline": True, "color": BLUE}),
        (
            ": al dar de alta a alguien, se fija en qué región / zona / unidad “vive”. "
            "Eso define su rol en el árbol orgánico.",
            {},
        ),
    ]
)
p_parts(
    [
        ("Capa B — Mapa", {"bold": True, "underline": True, "color": BLUE}),
        (
            ": a quién puede ver en el mapa. Solo por ",
            {},
        ),
        ("jerarquía hacia abajo y de lado (pares)", {"bold": True}),
        (
            ". Nadie ve hacia arriba (ej. unidad no ve ubicación de zona ni de región).",
            {},
        ),
    ]
)
p_parts(
    [
        ("Capa C — Canal / grupo (radio)", {"bold": True, "underline": True, "color": BLUE}),
        (
            ": sala PTT aparte. Crear un canal tiene su propia cascada. "
            "Hablar en un canal exige ser ",
            {},
        ),
        ("miembro", {"bold": True, "underline": True}),
        (" del canal (no basta “ser de la región”).", {}),
    ]
)
rule(
    "Pertenencia ≠ Mapa ≠ Canal. Las tres se relacionan, pero no son el mismo selector."
)

h1("2. Lo que este documento YA NO incluye")
p_parts(
    [
        (
            "Se elimina por completo la función de que administradores (u otros) "
            "puedan “ocultar ubicación”.",
            {"bold": True},
        )
    ],
    shade="FFE6E6",
)
rule(
    "Nadie oculta su ubicación. La única restricción es la jerarquía del mapa "
    "(abajo no ve arriba)."
)
warn(
    "Si el código aún permite locationShare=hidden o menús «Ocultar», hay que "
    "retirarlos / ignorarlos para alinear producto con este documento."
)

h1("3. Jerarquía de ejemplo")
p_parts(
    [
        ("Región IV R.M.", {"bold": True}),
        (" → ", {}),
        ("Zona 48/a. Z.M.", {"bold": True}),
        (" → ", {}),
        ("Unidad / organismo", {"bold": True}),
        (".", {}),
    ]
)

h1("4. Roles — Pertenencia (Capa A)")

h2("4.1 Administrador (sistema / root)")
bullet("Pertenencia:", "Global. No elige región/zona/unidad.")
rule("Alta sin cascada.")

h2("4.2 Administrador de región")
bullet("Pertenencia:", "Solo la región (ej. IV R.M.). Ahí termina el alta.")
proposal(
    "Punto 1 resuelto: en el alta del admin de región NO se pide zona ni «todas las zonas». "
    "Su adscripción es la región completa. La zona/unidad se eligen solo al CREAR un canal "
    "(Capa C), no al crear a la persona. Así ya no se confunde “dónde vive” con “qué canal arma”."
)
warn("Hoy el alta aún pide zona o «todas». Hay que cambiar la UI/backend para terminar en región.")

h2("4.3 Usuario de región")
bullet("Pertenencia:", "La región (ej. IV R.M.). Misma idea: termina en región.")
rule("No administra; no crea canales.")

h2("4.4 Administrador de zona")
bullet("Pertenencia:", "Región + una zona concreta. Sin opción «Todas las zonas».")
rule("Ahí termina el alta.")

h2("4.5 Usuario de zona")
bullet("Pertenencia:", "Región + una zona concreta.")
rule(
    "Rol que el Alcance original omitía y aquí queda formal: existe entre admin de zona "
    "y usuario de unidad. No crea canales."
)

h2("4.6 Administrador de unidad")
bullet("Pertenencia:", "Región → Zona → Unidad. Ahí termina.")

h2("4.7 Usuario de unidad")
bullet("Pertenencia:", "Región → Zona → Unidad. Ahí termina.")
rule("No crea canales.")

h1("5. Roles — Mapa (Capa B)")
p_parts(
    [
        ("Principio único: ", {"bold": True, "underline": True}),
        ("mirar hacia abajo y a pares; nunca hacia arriba.", {"bold": True}),
    ]
)

# small table
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
for i, h in enumerate(["Quién mira", "A quién ve en el mapa"]):
    hdr[i].text = ""
    rp = hdr[i].paragraphs[0]
    rr = rp.add_run(h)
    set_run(rr, bold=True, size=10, color=WHITE)
    shade_cell(hdr[i], "0B3D91")

map_rows = [
    ("Administrador (root)", "Todos."),
    (
        "Admin de región",
        "Todos en su región: admins/usuarios de región, zona y unidad de esa región.",
    ),
    (
        "Usuario de región",
        "Todos en su región (misma cobertura territorial que el admin de región, "
        "pero sin poder de administración). No ve otras regiones.",
    ),
    (
        "Admin de zona",
        "Su zona: admins/usuarios de zona y todas las unidades de esa zona. "
        "No ve región ni otras zonas.",
    ),
    (
        "Usuario de zona",
        "Pares de su zona (usuarios/admins de zona de esa zona). "
        "No ve región. (Unidades: ver nota abajo.)",
    ),
    ("Admin de unidad", "Solo su unidad (admin y usuarios de esa unidad)."),
    ("Usuario de unidad", "Solo pares de su unidad. No ve zona ni región."),
]
for a, b in map_rows:
    row = table.add_row().cells
    for i, val in enumerate((a, b)):
        row[i].text = ""
        rp = row[i].paragraphs[0]
        rr = rp.add_run(val)
        set_run(rr, bold=(i == 0), size=9, color=BLACK)

p_parts(
    [
        ("Nota usuario de zona / unidades: ", {"bold": True, "color": ORANGE}),
        (
            "propuesta: el usuario de zona ve solo pares de zona (como hoy la matriz). "
            "Si se necesita que también vea unidades de su zona, se documenta como excepción "
            "explícita; por defecto, no.",
            {"color": ORANGE},
        ),
    ]
)
rule(
    "Ejemplo: un usuario de unidad NUNCA ve la ubicación de un admin o usuario de zona/región."
)
warn(
    "Hoy el usuario de región NO ve zona/unidad en la matriz. Hay que ampliar la matriz "
    "para cumplir la fila de este documento."
)

h1("6. Roles — Canales / grupos (Capa C)")

h2("6.1 Quién puede crear canales")
bullet("Pueden crear:", "root, admin de región, admin de zona, admin de unidad.")
bullet("No pueden crear:", "usuario de región, usuario de zona, usuario de unidad.")

h2("6.2 Admin de región — cómo crea un canal (Punto 2 resuelto)")
p_parts(
    [
        (
            "Aquí se corrige la asimetría del Alcance original (el de región quedaba "
            "más limitado que el de zona).",
            {"italic": True},
        )
    ]
)
proposal(
    "Punto 2 integrado — misma lógica que el admin de zona, pero anclada a la región: "
    "(1) Elige región (fijada a la suya). "
    "(2) Elige «Todas las zonas» → canal de REGIÓN: puede agregar gente de toda la región. "
    "(3) Elige una zona concreta → entonces elige «Todas las unidades» de esa zona "
    "(canal de ZONA) O una unidad concreta (canal de UNIDAD). "
    "Así el admin de región SÍ puede armar un canal de toda una zona, igual que el admin "
    "de esa zona; no se le obliga a una sola unidad si quiere cobertura de zona."
)
rule(
    "Resumen canal admin de región: Todas las zonas = región · Zona + todas unidades = zona · "
    "Zona + una unidad = unidad."
)
p_parts(
    [
        ("Por qué se descarta la regla vieja “zona → solo unidad”: ", {"bold": True}),
        (
            "contradecía el sentido común (el jefe de región podía menos que el de zona) "
            "y mezclaba pertenencia con diseño del canal.",
            {},
        ),
    ]
)

h2("6.3 Admin de zona — cómo crea un canal")
bullet("Región y zona:", "bloqueadas a las suyas (solo se muestran).")
bullet("Unidad:", "«Todas las unidades» → canal de zona · o una unidad → canal de unidad.")
rule("Coincide con la propuesta del admin de región cuando este acota a una zona.")

h2("6.4 Admin de unidad — cómo crea un canal")
bullet("Todo bloqueado:", "solo su unidad. Nombra el canal y agrega miembros de esa unidad.")

h2("6.5 Miembros de un canal")
rule(
    "Solo se agregan personas cuya pertenencia caiga bajo el alcance del canal "
    "(región / zona / unidad) y que el creador tenga derecho a gestionar."
)
rule(
    "Para hablar (PTT) hay que ser miembro del canal. El mapa no da radio automático."
)
warn(
    "Confirmar en producto: listados de radio (app) usan membresía; no abrir canales "
    "solo por jerarquía salvo decisión futura explícita."
)

h1("7. Tabla única de verdad")

t2 = doc.add_table(rows=1, cols=4)
t2.style = "Table Grid"
hcells = t2.rows[0].cells
for i, h in enumerate(["Rol", "Pertenencia (alta)", "Mapa", "Crea canales"]):
    hcells[i].text = ""
    rp = hcells[i].paragraphs[0]
    rr = rp.add_run(h)
    set_run(rr, bold=True, size=9, color=WHITE)
    shade_cell(hcells[i], "0D6E6E")

truth = [
    ("Administrador", "Global", "Todos", "Cualquier alcance"),
    ("Admin región", "Solo región", "Toda su región ↓", "Región / zona / unidad (ver §6.2)"),
    ("Usuario región", "Solo región", "Toda su región ↓", "No"),
    ("Admin zona", "Región + zona", "Su zona ↓", "Zona o unidad"),
    ("Usuario zona", "Región + zona", "Pares de zona", "No"),
    ("Admin unidad", "Región+zona+unidad", "Su unidad", "Solo su unidad"),
    ("Usuario unidad", "Región+zona+unidad", "Pares de unidad", "No"),
]
for vals in truth:
    row = t2.add_row().cells
    for i, val in enumerate(vals):
        row[i].text = ""
        rp = row[i].paragraphs[0]
        rr = rp.add_run(val)
        set_run(rr, bold=(i == 0), size=8, color=BLACK)

h1("8. Contradicciones del Alcance 1/2 — cómo quedaron resueltas")

t3 = doc.add_table(rows=1, cols=3)
t3.style = "Table Grid"
for i, h in enumerate(["Antes (problema)", "Ahora (Alcance 3)", "Estado"]):
    c = t3.rows[0].cells[i]
    c.text = ""
    rp = c.paragraphs[0]
    rr = rp.add_run(h)
    set_run(rr, bold=True, size=9, color=WHITE)
    shade_cell(c, "0B3D91")

fixes = [
    (
        "Ocultar ubicación de admins vs quién igual los ve",
        "Se elimina ocultar. Solo jerarquía de mapa.",
        "Resuelto en documento",
    ),
    (
        "Admin región: alta pedía zona; grupos también — confusión",
        "Alta = solo región. Zona/unidad solo al crear canal (Punto 1).",
        "Resuelto en documento",
    ),
    (
        "Admin región con zona forzada a unidad (peor que admin zona)",
        "Misma lógica que admin zona: zona + todas unidades o una unidad (Punto 2).",
        "Resuelto en documento",
    ),
    (
        "Usuario de zona no existía en el Word",
        "Queda definido formalmente (§4.5 / §5).",
        "Resuelto en documento",
    ),
    (
        "Usuario región: mapa amplio vs radio",
        "Mapa amplio de región; radio solo por membresía (explícito).",
        "Aclarado (sin contradicción)",
    ),
]
for a, b, c in fixes:
    row = t3.add_row().cells
    for i, val in enumerate((a, b, c)):
        row[i].text = ""
        rp = row[i].paragraphs[0]
        color = GREEN if i == 2 else BLACK
        rr = rp.add_run(val)
        set_run(rr, bold=(i == 2), size=8, color=color)

h1("9. Checklist para implementar en el sistema")
warn("1) Alta admin/usuario de región: terminar en región (quitar zona obligatoria).")
warn(
    "2) Grupos admin de región: permitir zona + «todas las unidades» (canal zona), "
    "no solo unidad."
)
warn("3) Matriz mapa: usuario de región ve zona y unidad de su región.")
warn("4) Retirar UI/API de ocultar ubicación (locationShare hidden) para todos.")
warn("5) Mantener: unidad no ve zona/región; zona no ve región.")
rule(
    "Hasta cerrar el checklist en código, el documento manda; el software se alinea a esto."
)

h1("10. Glosario")
bullet("Pertenencia / adscripción:", "Dónde vive la persona en el árbol (Capa A).")
bullet("Mapa:", "Quién aparece en el mapa para cada rol (Capa B).")
bullet("Canal / grupo:", "Sala de radio PTT (Capa C).")
bullet("Miembro:", "Usuario agregado al canal; condición para hablar.")
bullet("Hacia abajo:", "Niveles inferiores en el árbol (región → zona → unidad).")
bullet("Hacia arriba:", "Niveles superiores; no se ven en el mapa.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— Fin · Reglas de pertenencia y canales (Alcance 3) —")
set_run(r, italic=True, size=9, color=GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Sustituye Alcance.docx y Alcance 2.docx como referencia de producto. "
    "Basado en revisión TacticalPtx · sept 2026."
)
set_run(r, italic=True, size=8, color=GRAY)

os.makedirs(r"C:\pulsanet_soporte\Documentos", exist_ok=True)
doc.save(OUT)
doc.save(OUT2)
print("OK", OUT)
print("OK", OUT2)
